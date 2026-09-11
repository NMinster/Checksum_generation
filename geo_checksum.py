#!/usr/bin/env python3
"""
geo_checksum.py - generate file checksums for a GEO submission and write them
into the GEO metadata spreadsheet.

Typical workflow
----------------
1. Compute checksums for every data file on the external drive, saving them to
   a CSV as you go (safe to interrupt and re-run; finished files are skipped):

       python geo_checksum.py scan /Volumes/MyDrive/project --out checksums.csv

2. Write those checksums into the GEO metadata Excel sheet, matching each
   "file name" cell in the RAW FILES / PROCESSED DATA FILES sections:

       python geo_checksum.py fill checksums.csv --excel GEO_metadata.xlsx

   Or do both in one step:

       python geo_checksum.py scan /Volumes/MyDrive/project --out checksums.csv --excel GEO_metadata.xlsx

Or let the tool find everything itself: it looks for an attached external
drive, finds the GEO metadata workbook stored on it, hashes the data files in
that folder and fills the workbook in place:

       python geo_checksum.py auto

GEO asks for MD5 checksums, which is the default algorithm.
"""

from __future__ import annotations

import argparse
import csv
import fnmatch
import hashlib
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

__version__ = "1.0.0"

# File types GEO typically wants checksums for. Used when --pattern is not given.
DEFAULT_PATTERNS = [
    "*.fastq", "*.fastq.gz", "*.fq", "*.fq.gz",
    "*.bam", "*.cram", "*.sra",
    "*.bw", "*.bigwig", "*.bigWig", "*.bedgraph", "*.bedGraph", "*.bedgraph.gz",
    "*.bed", "*.bed.gz", "*.narrowPeak", "*.broadPeak", "*.gtf", "*.gtf.gz",
    "*.h5", "*.h5ad", "*.mtx", "*.mtx.gz", "*.loom",
    "*.txt", "*.txt.gz", "*.tsv", "*.tsv.gz", "*.csv", "*.csv.gz", "*.xlsx",
    "*.tar", "*.tar.gz", "*.tgz", "*.zip",
]

# Files that are never data files and should not be hashed.
SKIP_NAMES = {".DS_Store", "Thumbs.db", "desktop.ini"}
SKIP_DIR_PREFIXES = (".", "$RECYCLE.BIN", "System Volume Information")

CSV_FIELDS = ["file_name", "checksum", "algorithm", "size_bytes", "mtime", "path"]

CHUNK_SIZE = 8 * 1024 * 1024  # 8 MiB reads suit slow external drives well.


# --------------------------------------------------------------------------- #
# Hashing
# --------------------------------------------------------------------------- #

@dataclass
class ChecksumRecord:
    file_name: str
    checksum: str
    algorithm: str
    size_bytes: int
    mtime: float
    path: str

    def as_row(self) -> dict:
        return {
            "file_name": self.file_name,
            "checksum": self.checksum,
            "algorithm": self.algorithm,
            "size_bytes": str(self.size_bytes),
            "mtime": repr(self.mtime),
            "path": self.path,
        }


def hash_file(path: Path, algorithm: str = "md5", progress=None) -> str:
    """Return the hex digest of *path*, reading it in large chunks."""
    h = hashlib.new(algorithm)
    with open(path, "rb") as fh:
        while True:
            chunk = fh.read(CHUNK_SIZE)
            if not chunk:
                break
            h.update(chunk)
            if progress is not None:
                progress(len(chunk))
    return h.hexdigest()


def iter_data_files(root: Path, patterns: Iterable[str], recursive: bool = True) -> Iterator[Path]:
    """Yield files under *root* whose name matches any of *patterns*."""
    patterns = list(patterns)
    if root.is_file():
        yield root
        return
    if not recursive:
        for p in sorted(root.iterdir()):
            if p.is_file() and _matches(p.name, patterns):
                yield p
        return
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith(SKIP_DIR_PREFIXES))
        for name in sorted(filenames):
            if _matches(name, patterns):
                yield Path(dirpath) / name


def _matches(name: str, patterns: list[str]) -> bool:
    if name in SKIP_NAMES or name.startswith("._") or name.startswith("~$"):
        return False
    if name.endswith(".bak.xlsx") or name == "checksums.csv":
        return False
    return any(fnmatch.fnmatch(name, pat) for pat in patterns)


# --------------------------------------------------------------------------- #
# CSV cache (lets a long scan be interrupted and resumed)
# --------------------------------------------------------------------------- #

def read_checksum_csv(path: Path) -> list[ChecksumRecord]:
    if not path.exists():
        return []
    records = []
    with open(path, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            try:
                records.append(ChecksumRecord(
                    file_name=row["file_name"],
                    checksum=row["checksum"],
                    algorithm=row.get("algorithm", "md5"),
                    size_bytes=int(row.get("size_bytes") or 0),
                    mtime=float(row.get("mtime") or 0),
                    path=row.get("path", ""),
                ))
            except (KeyError, ValueError) as exc:
                raise SystemExit(f"Could not parse {path}: bad row {row!r} ({exc})")
    return records


def write_checksum_csv(path: Path, records: Iterable[ChecksumRecord]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_FIELDS)
        writer.writeheader()
        for rec in records:
            writer.writerow(rec.as_row())


def write_md5sum_file(path: Path, records: Iterable[ChecksumRecord]) -> None:
    """Write a plain `md5sum -c` compatible file (one 'hash  filename' per line)."""
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        for rec in records:
            fh.write(f"{rec.checksum}  {rec.file_name}\n")


def _fmt_bytes(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024:
            return f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} PB"


def scan(root: Path, out_csv: Path | None, patterns: list[str], algorithm: str,
         recursive: bool = True, force: bool = False, quiet: bool = False,
         exclude: Iterable[Path] = ()) -> list[ChecksumRecord]:
    """Hash every matching file under *root*; reuse cached results from *out_csv*."""
    if not root.exists():
        raise SystemExit(f"Path not found: {root}")
    excluded = {str(Path(p).resolve()) for p in exclude}

    cached: dict[str, ChecksumRecord] = {}
    if out_csv is not None and not force:
        for rec in read_checksum_csv(out_csv):
            if rec.algorithm == algorithm:
                cached[rec.path] = rec

    files = [f for f in iter_data_files(root, patterns, recursive=recursive)
             if str(f.resolve()) not in excluded]
    if not files:
        print(f"No files matching {patterns} found under {root}", file=sys.stderr)
        return []

    todo, results = [], []
    for f in files:
        st = f.stat()
        key = str(f.resolve())
        rec = cached.get(key)
        if rec is not None and rec.size_bytes == st.st_size and abs(rec.mtime - st.st_mtime) < 1:
            results.append(rec)
        else:
            todo.append((f, st))

    total_bytes = sum(st.st_size for _, st in todo)
    if not quiet:
        print(f"Found {len(files)} file(s); {len(results)} already in cache, "
              f"{len(todo)} to hash ({_fmt_bytes(total_bytes)}).", file=sys.stderr)

    done_bytes = 0
    start = time.time()
    for i, (f, st) in enumerate(todo, 1):
        file_start = time.time()
        key = str(f.resolve())
        if not quiet:
            print(f"[{i}/{len(todo)}] {f.name} ({_fmt_bytes(st.st_size)}) ...",
                  file=sys.stderr, end="", flush=True)
        try:
            digest = hash_file(f, algorithm)
        except OSError as exc:
            print(f" ERROR: {exc}", file=sys.stderr)
            continue
        rec = ChecksumRecord(f.name, digest, algorithm, st.st_size, st.st_mtime, key)
        results.append(rec)
        done_bytes += st.st_size
        if out_csv is not None:
            # Save after every file so an interrupted run loses nothing.
            write_checksum_csv(out_csv, sorted(results, key=lambda r: r.path))
        if not quiet:
            elapsed = time.time() - start
            rate = done_bytes / elapsed if elapsed > 0 else 0
            remaining = (total_bytes - done_bytes) / rate if rate > 0 else 0
            print(f" {digest}  [{time.time() - file_start:.0f}s, "
                  f"{_fmt_bytes(rate)}/s, ~{remaining / 60:.0f} min left]", file=sys.stderr)

    results.sort(key=lambda r: r.path)
    dupes = _duplicate_names(results)
    if dupes:
        print("WARNING: the same file name appears in more than one folder. GEO matches on "
              "file name only, so these must be renamed to be unique:", file=sys.stderr)
        for name, paths in dupes.items():
            print(f"  {name}:", file=sys.stderr)
            for p in paths:
                print(f"    {p}", file=sys.stderr)
    return results


def _duplicate_names(records: list[ChecksumRecord]) -> dict[str, list[str]]:
    by_name: dict[str, list[str]] = {}
    for r in records:
        by_name.setdefault(r.file_name, []).append(r.path)
    return {n: p for n, p in by_name.items() if len(p) > 1}


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

FILE_NAME_HEADER = re.compile(r"^\s*file\s*name\s*$", re.IGNORECASE)
CHECKSUM_HEADER = re.compile(r"checksum|md5", re.IGNORECASE)


@dataclass
class Section:
    """A block in the sheet with a 'file name' column and a 'checksum' column."""
    header_row: int
    name_col: int
    checksum_col: int


def find_sections(ws) -> list[Section]:
    """Locate every header row that has both a 'file name' and a 'checksum' column.

    The GEO metadata template has two such blocks (RAW FILES and PROCESSED DATA
    FILES); a plain user-made sheet has one on row 1.
    """
    sections = []
    for row in ws.iter_rows():
        name_col = checksum_col = header_row = None
        for cell in row:
            v = cell.value
            if not isinstance(v, str):
                continue
            if name_col is None and FILE_NAME_HEADER.match(v):
                name_col, header_row = cell.column, cell.row
            elif checksum_col is None and CHECKSUM_HEADER.search(v):
                checksum_col = cell.column
        if name_col is not None and checksum_col is not None:
            sections.append(Section(header_row, name_col, checksum_col))
    return sections


def _col_index(ws, spec: str) -> int:
    """Turn a column letter ('C') or header text ('file checksum') into an index."""
    from openpyxl.utils import column_index_from_string
    if re.fullmatch(r"[A-Za-z]{1,3}", spec):
        return column_index_from_string(spec.upper())
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().lower() == spec.strip().lower():
                return cell.column
    raise SystemExit(f"Could not find a column named {spec!r} in sheet {ws.title!r}")


def fill_workbook(excel_path: Path, records: list[ChecksumRecord], sheet: str | None = None,
                  name_col: str | None = None, checksum_col: str | None = None,
                  output: Path | None = None, overwrite: bool = False,
                  backup: bool = True, quiet: bool = False) -> dict:
    """Write checksums next to matching file names in *excel_path*.

    Returns a summary dict with 'filled', 'unchanged', 'missing' (names in the
    sheet with no checksum) and 'unused' (checksums not referenced by the sheet).
    """
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl is required for Excel support: pip install openpyxl")

    if not excel_path.exists():
        raise SystemExit(f"Excel file not found: {excel_path}")

    by_name = {r.file_name: r for r in records}
    by_name_ci = {r.file_name.lower(): r for r in records}

    wb = openpyxl.load_workbook(excel_path)
    if sheet:
        if sheet not in wb.sheetnames:
            raise SystemExit(f"Sheet {sheet!r} not found. Available: {wb.sheetnames}")
        sheets = [wb[sheet]]
    else:
        sheets = wb.worksheets

    filled, unchanged, missing, used = [], [], [], set()
    for ws in sheets:
        if name_col and checksum_col:
            sections = [Section(0, _col_index(ws, name_col), _col_index(ws, checksum_col))]
        else:
            sections = find_sections(ws)
        if not sections:
            continue
        # Assign each data row to the nearest section header above it.
        sections.sort(key=lambda s: s.header_row)
        header_rows = {s.header_row for s in sections}
        for row_idx in range(1, ws.max_row + 1):
            section = None
            for s in sections:
                if s.header_row < row_idx:
                    section = s
            if section is None:
                continue
            # The GEO template puts a section title (e.g. "PROCESSED DATA FILES")
            # on the row directly above each header row; that is not a file name.
            if row_idx + 1 in header_rows:
                continue
            name_cell = ws.cell(row=row_idx, column=section.name_col)
            raw = name_cell.value
            if not isinstance(raw, str) or not raw.strip():
                continue
            name = raw.strip()
            if FILE_NAME_HEADER.match(name):
                continue
            # Allow the sheet to contain a path; GEO only needs the base name.
            base = Path(name.replace("\\", "/")).name
            rec = by_name.get(base) or by_name_ci.get(base.lower())
            if rec is None:
                missing.append((ws.title, row_idx, name))
                continue
            used.add(rec.file_name)
            target = ws.cell(row=row_idx, column=section.checksum_col)
            existing = target.value
            if existing not in (None, "") and str(existing).strip() and not overwrite:
                if str(existing).strip().lower() == rec.checksum.lower():
                    unchanged.append((ws.title, row_idx, name))
                else:
                    print(f"WARNING: {ws.title}!{target.coordinate} already has a different "
                          f"checksum for {name}; use --overwrite to replace it.", file=sys.stderr)
                    missing.append((ws.title, row_idx, name))
                continue
            target.value = rec.checksum
            filled.append((ws.title, row_idx, name))

    unused = sorted(set(by_name) - used)

    out = output or excel_path
    if out == excel_path and backup:
        bak = excel_path.with_name(excel_path.stem + ".bak" + excel_path.suffix)
        shutil.copy2(excel_path, bak)
        if not quiet:
            print(f"Backup saved to {bak}", file=sys.stderr)
    wb.save(out)

    if not quiet:
        print(f"Wrote {len(filled)} checksum(s) to {out}"
              + (f" ({len(unchanged)} already correct)" if unchanged else ""), file=sys.stderr)
        if missing:
            print(f"{len(missing)} file name(s) in the sheet had no checksum:", file=sys.stderr)
            for title, r, n in missing:
                print(f"  {title} row {r}: {n}", file=sys.stderr)
        if unused:
            print(f"{len(unused)} scanned file(s) are not listed in the sheet:", file=sys.stderr)
            for n in unused:
                print(f"  {n}", file=sys.stderr)
    return {"filled": filled, "unchanged": unchanged, "missing": missing, "unused": unused}



# --------------------------------------------------------------------------- #
# Auto mode: find the external drive and the GEO workbook on it
# --------------------------------------------------------------------------- #

WORKBOOK_SEARCH_DEPTH = 4  # how deep to look for the .xlsx on the drive


def candidate_drives() -> list[Path]:
    """Return mount points that look like external / removable drives."""
    drives: list[Path] = []
    if sys.platform == "darwin":
        root_dev = os.stat("/").st_dev
        for p in sorted(Path("/Volumes").glob("*")):
            try:
                if p.is_dir() and p.stat().st_dev != root_dev:
                    drives.append(p)
            except OSError:
                continue
    elif sys.platform.startswith("win"):
        import ctypes
        import string
        kernel32 = ctypes.windll.kernel32  # type: ignore[attr-defined]
        system_drive = os.environ.get("SystemDrive", "C:").rstrip("\\").upper()
        bitmask = kernel32.GetLogicalDrives()
        for i, letter in enumerate(string.ascii_uppercase):
            if not bitmask & (1 << i):
                continue
            root = f"{letter}:\\"
            # 2 = DRIVE_REMOVABLE, 3 = DRIVE_FIXED (USB hard drives usually report as fixed)
            if kernel32.GetDriveTypeW(root) in (2, 3) and f"{letter}:" != system_drive:
                drives.append(Path(root))
    else:
        user = os.environ.get("USER", "")
        for base in (f"/media/{user}", "/media", f"/run/media/{user}", "/mnt"):
            for p in sorted(Path(base).glob("*")):
                if p.is_dir() and p not in drives and not p.name.startswith("."):
                    drives.append(p)
    return drives


def find_geo_workbooks(root: Path, max_depth: int = WORKBOOK_SEARCH_DEPTH) -> list[Path]:
    """Find .xlsx files under *root* that contain a 'file name' / 'checksum' section."""
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl is required: pip install openpyxl")
    found = []
    root_depth = len(root.resolve().parts)
    for dirpath, dirnames, filenames in os.walk(root):
        depth = len(Path(dirpath).resolve().parts) - root_depth
        dirnames[:] = sorted(d for d in dirnames if not d.startswith(SKIP_DIR_PREFIXES))
        if depth >= max_depth:
            dirnames[:] = []
        for name in sorted(filenames):
            if not name.lower().endswith((".xlsx", ".xlsm")) or name.startswith(("~$", "._")):
                continue
            if name.lower().endswith(".bak.xlsx"):
                continue
            path = Path(dirpath) / name
            try:
                wb = openpyxl.load_workbook(path, read_only=True)
                is_geo = any(find_sections(ws) for ws in wb.worksheets)
                wb.close()
            except Exception:  # corrupt / locked / not really an xlsx
                continue
            if is_geo:
                found.append(path)
    return found


def _pick_workbook(workbooks: list[Path]) -> Path:
    if len(workbooks) == 1:
        return workbooks[0]
    preferred = [w for w in workbooks if re.search(r"geo|metadata|template", w.name, re.I)]
    if len(preferred) == 1:
        return preferred[0]
    listing = "\n".join(f"  {w}" for w in workbooks)
    raise SystemExit(f"Found more than one GEO-style workbook; pick one with --excel:\n{listing}")


def _confirm(prompt: str, assume_yes: bool) -> bool:
    if assume_yes or not sys.stdin.isatty():
        return True
    answer = input(f"{prompt} [Y/n] ").strip().lower()
    return answer in ("", "y", "yes")


def auto(drive: Path | None = None, excel: Path | None = None, root: Path | None = None,
         patterns: list[str] | None = None, algorithm: str = "md5", assume_yes: bool = False,
         overwrite: bool = False, quiet: bool = False) -> dict:
    """Locate drive -> workbook -> data folder, then scan and fill in one go."""
    # 1. Which drive?
    if drive is None and excel is None and root is None:
        drives = candidate_drives()
        if not drives:
            raise SystemExit("No external drive found. Plug it in, or pass --drive / --excel.")
        if len(drives) > 1:
            with_wb = [d for d in drives if find_geo_workbooks(d)]
            if len(with_wb) == 1:
                drives = with_wb
            else:
                listing = "\n".join(f"  {d}" for d in (with_wb or drives))
                raise SystemExit(f"More than one external drive is attached; pick one with "
                                 f"--drive:\n{listing}")
        drive = drives[0]
    if drive is not None and not drive.exists():
        raise SystemExit(f"Drive not found: {drive}")

    # 2. Which workbook?
    if excel is None:
        search_root = root or drive
        workbooks = find_geo_workbooks(search_root)
        if not workbooks:
            raise SystemExit(f"No GEO metadata workbook (.xlsx with 'file name' and 'file "
                             f"checksum' columns) found under {search_root}. Pass --excel.")
        excel = _pick_workbook(workbooks)
    if not excel.exists():
        raise SystemExit(f"Workbook not found: {excel}")

    # 3. Which folder holds the data files? Default: the folder the workbook lives in.
    data_root = root or excel.parent
    out_csv = excel.parent / "checksums.csv"

    print(f"Drive:     {drive or '(not auto-detected)'}", file=sys.stderr)
    print(f"Workbook:  {excel}", file=sys.stderr)
    print(f"Data root: {data_root}", file=sys.stderr)
    print(f"Checksums: {out_csv}", file=sys.stderr)
    if not _confirm("Hash every data file under the data root and fill the workbook?", assume_yes):
        raise SystemExit("Cancelled.")

    records = scan(data_root, out_csv, patterns or DEFAULT_PATTERNS, algorithm,
                   quiet=quiet, exclude=[excel, out_csv])
    if not records:
        raise SystemExit("No data files found; nothing written.")
    write_checksum_csv(out_csv, records)
    write_md5sum_file(excel.parent / "checksums.md5", records)
    summary = fill_workbook(excel, records, overwrite=overwrite, quiet=quiet)
    summary["workbook"] = excel
    summary["csv"] = out_csv
    return summary


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="geo_checksum",
        description="Generate MD5 checksums for GEO submission files and fill them into "
                    "the GEO metadata spreadsheet.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Typical workflow" + __doc__.split("Typical workflow")[1],
    )
    p.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    sub = p.add_subparsers(dest="command", required=True)

    excel_opts = argparse.ArgumentParser(add_help=False)
    excel_opts.add_argument("--excel", type=Path, help="GEO metadata .xlsx to update")
    excel_opts.add_argument("--sheet", help="Only update this sheet (default: all sheets)")
    excel_opts.add_argument("--name-col",
                            help="Column letter or header text holding file names "
                                 "(default: auto-detect 'file name' headers)")
    excel_opts.add_argument("--checksum-col",
                            help="Column letter or header text to write checksums into "
                                 "(default: auto-detect 'checksum' headers)")
    excel_opts.add_argument("--output", type=Path,
                            help="Save the updated workbook here instead of in place")
    excel_opts.add_argument("--overwrite", action="store_true",
                            help="Replace checksums already present in the sheet")
    excel_opts.add_argument("--no-backup", action="store_true",
                            help="Do not keep a .bak copy when editing in place")

    s = sub.add_parser("scan", parents=[excel_opts],
                       help="Compute checksums for files under a folder (or a single file)")
    s.add_argument("paths", nargs="+", type=Path, help="Folder(s) or file(s) on the external drive")
    s.add_argument("--out", type=Path, default=Path("checksums.csv"),
                   help="CSV to write results to; also used as a resume cache (default: checksums.csv)")
    s.add_argument("--md5sum-file", type=Path,
                   help="Also write a plain 'md5sum -c' style text file")
    s.add_argument("--pattern", action="append",
                   help="Glob for files to include, e.g. '*.fastq.gz'. Repeatable. "
                        "Use '*' for everything. Default: common sequencing/processed file types.")
    s.add_argument("--algorithm", default="md5", choices=sorted(hashlib.algorithms_guaranteed),
                   help="Hash algorithm (GEO uses md5; default: md5)")
    s.add_argument("--no-recursive", action="store_true", help="Do not descend into subfolders")
    s.add_argument("--force", action="store_true", help="Re-hash files even if cached in --out")
    s.add_argument("-q", "--quiet", action="store_true")

    f = sub.add_parser("fill", parents=[excel_opts],
                       help="Write checksums from a CSV (made by 'scan') into an Excel sheet")
    f.add_argument("csv", type=Path, help="checksums.csv produced by 'scan'")
    f.add_argument("-q", "--quiet", action="store_true")

    a = sub.add_parser("auto",
                       help="Find the external drive and the GEO workbook on it, then scan and fill")
    a.add_argument("--drive", type=Path, help="Mount point of the drive (default: auto-detect)")
    a.add_argument("--excel", type=Path, help="Workbook to fill (default: search the drive)")
    a.add_argument("--root", type=Path,
                   help="Folder holding the data files (default: the workbook's folder)")
    a.add_argument("--pattern", action="append", help="Glob for files to include (repeatable)")
    a.add_argument("--algorithm", default="md5", choices=sorted(hashlib.algorithms_guaranteed))
    a.add_argument("--overwrite", action="store_true",
                   help="Replace checksums already present in the sheet")
    a.add_argument("-y", "--yes", action="store_true", help="Do not ask for confirmation")
    a.add_argument("-q", "--quiet", action="store_true")

    v = sub.add_parser("verify", help="Re-hash files and compare against a checksums CSV")
    v.add_argument("csv", type=Path, help="checksums.csv produced by 'scan'")
    v.add_argument("-q", "--quiet", action="store_true")
    return p


def _excel_kwargs(args) -> dict:
    return dict(sheet=args.sheet, name_col=args.name_col, checksum_col=args.checksum_col,
                output=args.output, overwrite=args.overwrite, backup=not args.no_backup,
                quiet=args.quiet)


def cmd_scan(args) -> int:
    patterns = args.pattern or DEFAULT_PATTERNS
    records: list[ChecksumRecord] = []
    for path in args.paths:
        records.extend(scan(path, args.out, patterns, args.algorithm,
                            recursive=not args.no_recursive, force=args.force, quiet=args.quiet))
    if not records:
        return 1
    records.sort(key=lambda r: r.path)
    write_checksum_csv(args.out, records)
    if not args.quiet:
        print(f"Saved {len(records)} checksum(s) to {args.out}", file=sys.stderr)
    if args.md5sum_file:
        write_md5sum_file(args.md5sum_file, records)
    if args.excel:
        if args.excel.resolve() == args.out.resolve():
            raise SystemExit("--excel and --out must be different files")
        fill_workbook(args.excel, records, **_excel_kwargs(args))
    if not args.quiet and not args.excel:
        print(f"Next: python {Path(sys.argv[0]).name} fill {args.out} --excel <GEO_metadata.xlsx>",
              file=sys.stderr)
    return 0


def cmd_fill(args) -> int:
    if not args.excel:
        raise SystemExit("--excel is required for 'fill'")
    records = read_checksum_csv(args.csv)
    if not records:
        raise SystemExit(f"No checksums found in {args.csv}")
    summary = fill_workbook(args.excel, records, **_excel_kwargs(args))
    return 0 if not summary["missing"] else 2


def cmd_auto(args) -> int:
    summary = auto(drive=args.drive, excel=args.excel, root=args.root, patterns=args.pattern,
                   algorithm=args.algorithm, assume_yes=args.yes, overwrite=args.overwrite,
                   quiet=args.quiet)
    print(f"Done. {len(summary['filled'])} checksum(s) written to {summary['workbook']}",
          file=sys.stderr)
    return 0 if not summary["missing"] else 2


def cmd_verify(args) -> int:
    records = read_checksum_csv(args.csv)
    bad = 0
    for rec in records:
        p = Path(rec.path)
        if not p.exists():
            print(f"MISSING  {rec.file_name}  ({rec.path})")
            bad += 1
            continue
        digest = hash_file(p, rec.algorithm)
        if digest == rec.checksum:
            if not args.quiet:
                print(f"OK       {rec.file_name}")
        else:
            print(f"CHANGED  {rec.file_name}  expected {rec.checksum} got {digest}")
            bad += 1
    print(f"{len(records) - bad}/{len(records)} file(s) verified OK", file=sys.stderr)
    return 1 if bad else 0


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        return {"scan": cmd_scan, "fill": cmd_fill, "verify": cmd_verify,
                "auto": cmd_auto}[args.command](args)
    except KeyboardInterrupt:
        print("\nInterrupted. Progress so far is saved in the CSV; re-run to resume.",
              file=sys.stderr)
        return 130


if __name__ == "__main__":
    sys.exit(main())
