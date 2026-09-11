import hashlib
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import geo_checksum as gc  # noqa: E402


def md5(data: bytes) -> str:
    return hashlib.md5(data).hexdigest()


@pytest.fixture
def drive(tmp_path):
    """A fake external drive with raw and processed files in subfolders."""
    raw = tmp_path / "drive" / "raw"
    proc = tmp_path / "drive" / "processed"
    raw.mkdir(parents=True)
    proc.mkdir(parents=True)
    files = {
        raw / "sampleA_R1.fastq.gz": b"AAAA" * 1000,
        raw / "sampleA_R2.fastq.gz": b"CCCC" * 1000,
        proc / "counts.txt": b"gene\tcount\nX\t1\n",
        raw / ".DS_Store": b"junk",
        raw / "._sampleA_R1.fastq.gz": b"applejunk",
        raw / "notes.docx": b"not a data file",
    }
    for p, data in files.items():
        p.write_bytes(data)
    return tmp_path / "drive", files


@pytest.fixture
def geo_workbook(tmp_path):
    """Mimic the GEO metadata template: two sections with their own header rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metadata"
    rows = [
        ["SERIES"],
        ["title", "My study"],
        [],
        ["RAW FILES"],
        ["file name", "file type", "file checksum", "instrument model"],
        ["sampleA_R1.fastq.gz", "fastq", None, "NovaSeq"],
        ["sampleA_R2.fastq.gz", "fastq", None, "NovaSeq"],
        ["missing_R1.fastq.gz", "fastq", None, "NovaSeq"],
        [],
        ["PROCESSED DATA FILES"],
        ["file name", "file type", "file checksum"],
        ["counts.txt", "txt", None],
    ]
    for r in rows:
        ws.append(r)
    path = tmp_path / "GEO_metadata.xlsx"
    wb.save(path)
    return path


def test_scan_hashes_only_data_files(drive, tmp_path):
    root, files = drive
    out = tmp_path / "checksums.csv"
    records = gc.scan(root, out, gc.DEFAULT_PATTERNS, "md5", quiet=True)
    names = {r.file_name for r in records}
    assert names == {"sampleA_R1.fastq.gz", "sampleA_R2.fastq.gz", "counts.txt"}
    by_name = {r.file_name: r for r in records}
    assert by_name["sampleA_R1.fastq.gz"].checksum == md5(b"AAAA" * 1000)
    assert by_name["counts.txt"].checksum == md5(b"gene\tcount\nX\t1\n")
    # CSV round-trips
    assert {r.file_name for r in gc.read_checksum_csv(out)} == names


def test_scan_resumes_from_cache(drive, tmp_path, monkeypatch):
    root, _ = drive
    out = tmp_path / "checksums.csv"
    gc.scan(root, out, gc.DEFAULT_PATTERNS, "md5", quiet=True)

    calls = []
    real = gc.hash_file
    monkeypatch.setattr(gc, "hash_file", lambda p, a, progress=None: calls.append(p) or real(p, a))
    gc.scan(root, out, gc.DEFAULT_PATTERNS, "md5", quiet=True)
    assert calls == []  # everything served from cache

    # Modifying a file invalidates just that entry
    target = root / "processed" / "counts.txt"
    target.write_bytes(b"different")
    gc.scan(root, out, gc.DEFAULT_PATTERNS, "md5", quiet=True)
    assert [p.name for p in calls] == ["counts.txt"]

    # --force re-hashes everything
    calls.clear()
    gc.scan(root, out, gc.DEFAULT_PATTERNS, "md5", force=True, quiet=True)
    assert len(calls) == 3


def test_custom_pattern(drive, tmp_path):
    root, _ = drive
    records = gc.scan(root, None, ["*.docx"], "md5", quiet=True)
    assert [r.file_name for r in records] == ["notes.docx"]


def test_fill_geo_template_sections(drive, geo_workbook, tmp_path):
    root, _ = drive
    records = gc.scan(root, None, gc.DEFAULT_PATTERNS, "md5", quiet=True)
    summary = gc.fill_workbook(geo_workbook, records, quiet=True)

    assert {n for _, _, n in summary["filled"]} == {
        "sampleA_R1.fastq.gz", "sampleA_R2.fastq.gz", "counts.txt"}
    assert [n for _, _, n in summary["missing"]] == ["missing_R1.fastq.gz"]
    assert summary["unused"] == []

    ws = openpyxl.load_workbook(geo_workbook)["Metadata"]
    assert ws["C6"].value == md5(b"AAAA" * 1000)
    assert ws["C7"].value == md5(b"CCCC" * 1000)
    assert ws["C8"].value is None
    assert ws["C12"].value == md5(b"gene\tcount\nX\t1\n")
    # header rows untouched
    assert ws["C5"].value == "file checksum"
    assert ws["C11"].value == "file checksum"
    # backup written
    assert (tmp_path / "GEO_metadata.bak.xlsx").exists()


def test_fill_does_not_overwrite_without_flag(drive, geo_workbook):
    root, _ = drive
    wb = openpyxl.load_workbook(geo_workbook)
    wb["Metadata"]["C6"].value = "deadbeef"
    wb.save(geo_workbook)

    records = gc.scan(root, None, gc.DEFAULT_PATTERNS, "md5", quiet=True)
    summary = gc.fill_workbook(geo_workbook, records, quiet=True, backup=False)
    assert ("Metadata", 6, "sampleA_R1.fastq.gz") in summary["missing"]
    assert openpyxl.load_workbook(geo_workbook)["Metadata"]["C6"].value == "deadbeef"

    gc.fill_workbook(geo_workbook, records, quiet=True, backup=False, overwrite=True)
    assert openpyxl.load_workbook(geo_workbook)["Metadata"]["C6"].value == md5(b"AAAA" * 1000)


def test_fill_explicit_columns_and_output(drive, tmp_path):
    root, _ = drive
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sample", "Filename", "MD5"])
    ws.append(["A", "raw/sampleA_R1.fastq.gz", None])
    ws.append(["B", "SAMPLEA_R2.FASTQ.GZ", None])
    src = tmp_path / "sheet.xlsx"
    wb.save(src)

    records = gc.scan(root, None, gc.DEFAULT_PATTERNS, "md5", quiet=True)
    out = tmp_path / "sheet_filled.xlsx"
    summary = gc.fill_workbook(src, records, name_col="Filename", checksum_col="C",
                               output=out, quiet=True)
    assert len(summary["filled"]) == 2
    ws2 = openpyxl.load_workbook(out).active
    assert ws2["C2"].value == md5(b"AAAA" * 1000)   # path in cell -> basename match
    assert ws2["C3"].value == md5(b"CCCC" * 1000)   # case-insensitive match
    assert openpyxl.load_workbook(src).active["C2"].value is None  # source untouched


def test_cli_end_to_end(drive, geo_workbook, tmp_path):
    root, _ = drive
    out = tmp_path / "checksums.csv"
    md5file = tmp_path / "checksums.md5"
    rc = gc.main(["scan", str(root), "--out", str(out), "--md5sum-file", str(md5file),
                  "--excel", str(geo_workbook), "--no-backup", "-q"])
    assert rc == 0
    ws = openpyxl.load_workbook(geo_workbook)["Metadata"]
    assert ws["C12"].value == md5(b"gene\tcount\nX\t1\n")
    lines = md5file.read_text().splitlines()
    assert f"{md5(b'AAAA' * 1000)}  sampleA_R1.fastq.gz" in lines

    assert gc.main(["verify", str(out), "-q"]) == 0
    (root / "raw" / "sampleA_R2.fastq.gz").write_bytes(b"changed")
    assert gc.main(["verify", str(out), "-q"]) == 1

    # fill returns 2 when the sheet lists a file we have no checksum for
    assert gc.main(["fill", str(out), "--excel", str(geo_workbook), "--no-backup", "-q"]) == 2


# --------------------------------------------------------------------------- #
# auto mode
# --------------------------------------------------------------------------- #

@pytest.fixture
def fake_drive(tmp_path):
    """An 'external drive' with the GEO workbook sitting next to the data files."""
    drive = tmp_path / "Volumes" / "SEQDATA"
    proj = drive / "GEO_submission"
    (proj / "fastq").mkdir(parents=True)
    (proj / "fastq" / "s1_R1.fastq.gz").write_bytes(b"AAAA" * 10)
    (proj / "fastq" / "s1_R2.fastq.gz").write_bytes(b"CCCC" * 10)
    (proj / "counts.txt").write_bytes(b"x")
    (drive / "unrelated.xlsx").write_bytes(b"not a workbook")  # must be ignored

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["RAW FILES"])
    ws.append(["file name", "file type", "file checksum"])
    ws.append(["s1_R1.fastq.gz", "fastq", None])
    ws.append(["s1_R2.fastq.gz", "fastq", None])
    ws.append([])
    ws.append(["PROCESSED DATA FILES"])
    ws.append(["file name", "file type", "file checksum"])
    ws.append(["counts.txt", "txt", None])
    wb.save(proj / "GEO_metadata.xlsx")

    wb2 = openpyxl.Workbook()
    wb2.active.append(["Sample", "Notes"])
    wb2.save(proj / "lab_notes.xlsx")  # xlsx without checksum columns: ignored
    return drive


def test_find_geo_workbooks(fake_drive):
    found = gc.find_geo_workbooks(fake_drive)
    assert [p.name for p in found] == ["GEO_metadata.xlsx"]


def test_auto_detects_drive_and_fills_workbook(fake_drive, monkeypatch):
    monkeypatch.setattr(gc, "candidate_drives", lambda: [fake_drive])
    summary = gc.auto(assume_yes=True, quiet=True)
    wb_path = fake_drive / "GEO_submission" / "GEO_metadata.xlsx"
    assert summary["workbook"] == wb_path
    assert len(summary["filled"]) == 3
    assert summary["missing"] == []
    # The GEO workbook itself is never hashed; other spreadsheets are (they may be
    # processed data) and are merely reported as not listed in the sheet.
    assert summary["unused"] == ["lab_notes.xlsx"]
    ws = openpyxl.load_workbook(wb_path).active
    assert ws["C3"].value == md5(b"AAAA" * 10)
    assert ws["C8"].value == md5(b"x")
    assert (wb_path.parent / "checksums.csv").exists()
    assert (wb_path.parent / "checksums.md5").exists()
    assert (wb_path.parent / "GEO_metadata.bak.xlsx").exists()


def test_auto_no_drive(monkeypatch):
    monkeypatch.setattr(gc, "candidate_drives", lambda: [])
    with pytest.raises(SystemExit, match="No external drive"):
        gc.auto(assume_yes=True, quiet=True)


def test_auto_two_drives_picks_the_one_with_workbook(fake_drive, tmp_path, monkeypatch):
    other = tmp_path / "Volumes" / "BACKUP"
    other.mkdir()
    monkeypatch.setattr(gc, "candidate_drives", lambda: [other, fake_drive])
    summary = gc.auto(assume_yes=True, quiet=True)
    assert summary["workbook"].parent.parent == fake_drive


def test_auto_explicit_excel_skips_detection(fake_drive, monkeypatch):
    monkeypatch.setattr(gc, "candidate_drives", lambda: (_ for _ in ()).throw(AssertionError))
    wb_path = fake_drive / "GEO_submission" / "GEO_metadata.xlsx"
    rc = gc.main(["auto", "--excel", str(wb_path), "-y", "-q"])
    assert rc == 0


def test_candidate_drives_runs_on_this_platform():
    # Just make sure the platform-specific probing does not blow up.
    assert isinstance(gc.candidate_drives(), list)
